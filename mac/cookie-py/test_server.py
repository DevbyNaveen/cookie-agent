"""Offline tests for server.py. Run with the bundled interpreter:

    vendor/python/bin/python3 -I -B -m unittest discover -s cookie-py -v

No test touches the network: fetching goes through a fake resolver and a fake
transport, except the size-cap test, which talks to a local server directly.
"""

import gzip
import hashlib
import http.server
import json
import os
import subprocess
import sys
import tempfile
import threading
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import server  # noqa: E402
from server import ToolError  # noqa: E402


def resolver(table):
    def resolve(host, port):
        if host not in table:
            raise OSError(f"no such host {host}")
        return table[host]
    return resolve


class FakeWeb:
    """A transport that answers from a dict {(host, target): (status, headers, body)} and records every hop."""

    def __init__(self, pages):
        self.pages, self.hops = pages, []

    def __call__(self, scheme, host, port, ip, target, headers, deadline):
        self.hops.append((scheme, host, ip, target))
        assert headers["User-Agent"] == server.USER_AGENT
        return self.pages[(host, target)]


PUBLIC = "93.184.215.14"
HTML = (
    "<html><head><title>Field  notes</title></head><body>"
    "<nav>Home | About</nav>"
    "<article><h1>Report</h1><p>The river rose <b>two</b> <i>metres</i> overnight. Nobody was hurt.</p>"
    "<p>Crews\nworked&nbsp;until   dawn. " + "Filler sentence here. " * 20 + "</p></article>"
    "<script>var secret = 1;</script><footer>Copyright</footer></body></html>"
)


class SsrfGuard(unittest.TestCase):
    def test_every_blocked_range_is_refused(self):
        for ip in ["127.0.0.1", "127.9.9.9", "10.1.2.3", "172.16.0.1", "172.31.255.255", "192.168.1.1",
                   "169.254.169.254", "169.254.1.1", "100.64.0.1", "100.127.255.255", "0.0.0.0",
                   "224.0.0.1", "239.255.255.250", "::1", "::", "fd00::1", "fc00::1", "fe80::1", "ff02::1",
                   "::ffff:127.0.0.1", "::ffff:10.0.0.1", "64:ff9b::a9fe:a9fe", "2002:a00:1::1",
                   "192.0.2.1", "198.18.0.1", "240.0.0.1"]:
            with self.subTest(ip=ip):
                self.assertIsNotNone(server.refusal(ip), ip)

    def test_public_addresses_pass(self):
        for ip in [PUBLIC, "1.1.1.1", "2606:4700:4700::1111", "::ffff:1.1.1.1"]:
            with self.subTest(ip=ip):
                self.assertIsNone(server.refusal(ip))

    def test_a_host_resolving_to_private_is_refused(self):
        with self.assertRaisesRegex(ToolError, "non-public"):
            server.vet("https://intranet.example/", resolver({"intranet.example": ["10.0.0.7"]}))

    def test_one_private_answer_among_public_ones_is_refused(self):
        with self.assertRaisesRegex(ToolError, "non-public"):
            server.vet("http://mixed.example/", resolver({"mixed.example": [PUBLIC, "127.0.0.1"]}))

    def test_literal_and_odd_forms_go_through_resolution(self):
        for url, answer in [("http://169.254.169.254/latest/meta-data/", "169.254.169.254"),
                            ("http://2130706433/", "127.0.0.1"), ("http://[::1]:8080/", "::1")]:
            with self.subTest(url=url):
                host = server.urlsplit(url).hostname
                with self.assertRaises(ToolError):
                    server.vet(url, resolver({host: [answer]}))

    def test_only_http_and_https_without_credentials(self):
        for url in ["file:///etc/passwd", "ftp://example.com/", "gopher://x/", "example.com", "javascript:alert(1)",
                    "http://user:pw@example.com/"]:
            with self.subTest(url=url), self.assertRaises(ToolError):
                server.vet(url, resolver({"example.com": [PUBLIC]}))

    def test_a_redirect_to_a_private_address_is_refused_before_connecting(self):
        web = FakeWeb({("news.example", "/a"): (302, {"location": "http://internal.example/admin"}, b"")})
        resolve = resolver({"news.example": [PUBLIC], "internal.example": ["192.168.0.10"]})
        with self.assertRaisesRegex(ToolError, "non-public"):
            server.fetch_raw("https://news.example/a", resolve, web)
        self.assertEqual(len(web.hops), 1, "the private hop must never be requested")

    def test_a_redirect_to_the_metadata_address_is_refused(self):
        web = FakeWeb({("news.example", "/a"): (301, {"location": "http://169.254.169.254/latest/"}, b"")})
        resolve = resolver({"news.example": [PUBLIC], "169.254.169.254": ["169.254.169.254"]})
        with self.assertRaises(ToolError):
            server.fetch_raw("https://news.example/a", resolve, web)
        self.assertEqual(len(web.hops), 1)

    def test_a_redirect_to_another_scheme_is_refused(self):
        web = FakeWeb({("news.example", "/a"): (302, {"location": "file:///etc/passwd"}, b"")})
        with self.assertRaisesRegex(ToolError, "only http and https"):
            server.fetch_raw("https://news.example/a", resolver({"news.example": [PUBLIC]}), web)

    def test_redirects_are_followed_to_five_hops_and_no_further(self):
        pages = {("a.example", f"/{i}"): (302, {"location": f"/{i + 1}"}, b"") for i in range(6)}
        pages[("a.example", "/5")] = (200, {"content-type": "text/plain"}, b"arrived")
        final, _, body = server.fetch_raw("http://a.example/0", resolver({"a.example": [PUBLIC]}), FakeWeb(pages))
        self.assertEqual((final, body), ("http://a.example/5", b"arrived"))
        pages[("a.example", "/5")] = (302, {"location": "/6"}, b"")
        with self.assertRaisesRegex(ToolError, "more than 5 redirects"):
            server.fetch_raw("http://a.example/0", resolver({"a.example": [PUBLIC]}), FakeWeb(pages))

    def test_the_vetted_ip_is_the_one_connected_to(self):
        web = FakeWeb({("a.example", "/"): (200, {"content-type": "text/plain"}, b"ok")})
        server.fetch_raw("http://a.example/", resolver({"a.example": [PUBLIC]}), web)
        self.assertEqual(web.hops[0][2], PUBLIC)

    def test_a_compression_bomb_is_refused(self):
        bomb = gzip.compress(b"\0" * (server.MAX_BYTES + 10))
        with self.assertRaisesRegex(ToolError, "larger than"):
            server.decompress(bomb, "gzip")


class SizeCap(unittest.TestCase):
    """The real transport against a local server (called directly: vet would refuse loopback)."""

    def serve(self, body, chunked):
        class Handler(http.server.BaseHTTPRequestHandler):
            protocol_version = "HTTP/1.1"

            def do_GET(self):
                try:
                    self.reply()
                except ConnectionError:  # the client hung up at the cap, as it should
                    pass

            def reply(self):
                self.send_response(200)
                self.send_header("Content-Type", "text/html")
                if chunked:
                    self.send_header("Transfer-Encoding", "chunked")
                    self.end_headers()
                    for i in range(0, len(body), 65536):
                        part = body[i:i + 65536]
                        self.wfile.write(b"%x\r\n%s\r\n" % (len(part), part))
                    self.wfile.write(b"0\r\n\r\n")
                else:
                    self.send_header("Content-Length", str(len(body)))
                    self.send_header("Connection", "close")  # the response then owns (and closes) the socket
                    self.end_headers()
                    self.wfile.write(body)

            def log_message(self, *a):
                pass

        httpd = http.server.HTTPServer(("127.0.0.1", 0), Handler)
        threading.Thread(target=httpd.serve_forever, daemon=True).start()
        self.addCleanup(httpd.server_close)
        self.addCleanup(httpd.shutdown)
        return httpd.server_address[1]

    def get(self, port):
        deadline = server.time.monotonic() + 10
        return server.http_transport("http", "localhost", port, "127.0.0.1", "/", {"User-Agent": "t"}, deadline)

    def test_declared_and_streamed_oversize_bodies_are_refused(self):
        big = b"x" * (server.MAX_BYTES + 1)
        for chunked in (False, True):
            with self.subTest(chunked=chunked), self.assertRaisesRegex(ToolError, "larger than 5 MB"):
                self.get(self.serve(big, chunked))

    def test_a_body_under_the_cap_is_read_whole(self):
        for chunked in (False, True):
            with self.subTest(chunked=chunked):
                status, head, body = self.get(self.serve(b"<p>hi</p>", chunked))
                self.assertEqual((status, body), (200, b"<p>hi</p>"))


class Pages(unittest.TestCase):
    def setUp(self):
        server._CACHE.clear()
        self.web = FakeWeb({("news.example", "/r"): (200, {"content-type": "text/html; charset=utf-8",
                                                           "content-encoding": "gzip"}, gzip.compress(HTML.encode()))})
        self.resolve = resolver({"news.example": [PUBLIC]})

    def test_fetch_page_returns_title_main_text_and_hash(self):
        out = server.tool_fetch_page({"url": "https://news.example/r"}, self.resolve, self.web)
        self.assertEqual(out["title"], "Field notes")
        self.assertIn("The river rose two metres overnight.", out["text"])
        for absent in ("Home | About", "secret", "Copyright"):
            self.assertNotIn(absent, out["text"])
        self.assertEqual(out["sha256"], hashlib.sha256(out["text"].encode()).hexdigest())
        self.assertFalse(out["truncated"])
        self.assertTrue(out["fetched_at"].endswith("Z"))

    def test_truncation_is_said(self):
        out = server.tool_fetch_page({"url": "https://news.example/r", "max_chars": 20}, self.resolve, self.web)
        self.assertEqual(len(out["text"]), 20)
        self.assertTrue(out["truncated"])
        self.assertIn("truncated", out["note"])

    def test_a_page_without_a_main_region_falls_back_to_visible_text(self):
        web = FakeWeb({("a.example", "/"): (200, {"content-type": "text/html"},
                                            b"<html><body><h1>Hi</h1><div>Plain <a href=x>page</a>.</div>"
                                            b"<style>p{}</style></body></html>")})
        out = server.tool_fetch_page({"url": "http://a.example/"}, resolver({"a.example": [PUBLIC]}), web)
        self.assertEqual(out["title"], "Hi")
        self.assertIn("Plain page.", out["text"])
        self.assertNotIn("p{}", out["text"])

    def test_non_pages_and_http_errors_are_errors(self):
        web = FakeWeb({("a.example", "/pdf"): (200, {"content-type": "application/pdf"}, b"%PDF"),
                       ("a.example", "/gone"): (404, {}, b"")})
        for path, msg in (("/pdf", "not a web page"), ("/gone", "HTTP 404")):
            with self.subTest(path=path), self.assertRaisesRegex(ToolError, msg):
                server.tool_fetch_page({"url": f"http://a.example{path}"}, resolver({"a.example": [PUBLIC]}), web)

    def test_fetch_pages_reports_each_url(self):
        out = server.tool_fetch_pages({"urls": ["https://news.example/r", "http://10.0.0.1/"]},
                                      resolver({"news.example": [PUBLIC], "10.0.0.1": ["10.0.0.1"]}), self.web)
        first, second = out["results"]
        self.assertTrue(first["ok"])
        self.assertEqual((second["ok"], second["url"]), (False, "http://10.0.0.1/"))
        self.assertIn("non-public", second["error"])
        with self.assertRaises(ToolError):
            server.tool_fetch_pages({"urls": ["http://a/"] * 6})


class Quotes(unittest.TestCase):
    TEXT = "Intro line.\nThe river rose two\n metres  overnight. Nobody was hurt. Crews worked."

    def test_whitespace_is_normalised(self):
        hit = server.find_quote(self.TEXT, "river rose two metres\tovernight")
        self.assertTrue(hit["found"])
        self.assertEqual(hit["sentence"], "The river rose two metres overnight.")
        self.assertTrue(server.find_quote(self.TEXT, "Nobody was hurt.")["found"])

    def test_case_and_wording_must_match(self):
        self.assertFalse(server.find_quote(self.TEXT, "the river rose")["found"])
        self.assertFalse(server.find_quote(self.TEXT, "river rose three metres")["found"])

    def test_a_quote_across_sentences_returns_both(self):
        hit = server.find_quote(self.TEXT, "overnight. Nobody")
        self.assertEqual(hit["sentence"], "The river rose two metres overnight. Nobody was hurt.")

    def test_an_empty_quote_is_an_error(self):
        with self.assertRaises(ToolError):
            server.find_quote(self.TEXT, " \n ")

    def test_check_quote_checks_the_text_fetch_page_returned(self):
        server._CACHE.clear()
        web = FakeWeb({("news.example", "/r"): (200, {"content-type": "text/html"}, HTML.encode())})
        resolve = resolver({"news.example": [PUBLIC]})
        server.tool_fetch_page({"url": "https://news.example/r"}, resolve, web)
        out = server.tool_check_quote({"url": "https://news.example/r", "quote": "Crews worked until dawn."},
                                      resolve, web)
        self.assertTrue(out["found"] and out["from_cache"])
        self.assertEqual(len(web.hops), 1)
        server._CACHE.clear()
        out = server.tool_check_quote({"url": "https://news.example/r", "quote": "rose two metres"}, resolve, web)
        self.assertTrue(out["found"])
        self.assertFalse(out["from_cache"])


DECK = {
    "title": "Q3 review",
    "subtitle": "Prepared by Cookie",
    "slides": [
        {"title": "Highlights", "bullets": ["Revenue up 12%", "Two new markets"], "notes": "Say it slowly."},
        {"title": "Revenue", "chart": {"type": "column", "categories": ["Q1", "Q2", "Q3"],
                                       "series": [{"name": "2025", "values": [1, 2, 3]},
                                                  {"name": "2026", "values": [2, 3, 4.5]}]}},
        {"title": "Mix", "bullets": ["Mostly EU"], "chart": {"type": "pie", "categories": ["EU", "US"],
                                                             "series": [{"name": "Share", "values": [70, 30]}]}},
        {"title": "Team", "table": {"header": ["Name", "Role"], "rows": [["Ann", "Lead"], ["Bo", "Eng"]]}},
    ],
}


class Decks(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.root = os.path.realpath(tmp.name)
        self.roots = [os.path.join(self.root, "allowed")]
        os.mkdir(self.roots[0])

    def deck(self, path, **extra):
        return server.tool_make_deck({**DECK, "path": path, **extra}, self.roots)

    def test_decks_are_widescreen(self):
        from pptx import Presentation
        from pptx.util import Inches

        path = os.path.join(self.roots[0], "wide.pptx")
        self.deck(path)
        prs = Presentation(path)
        self.assertEqual((prs.slide_width, prs.slide_height), (Inches(13.333), Inches(7.5)))

    def test_deck_round_trips_through_python_pptx(self):
        from pptx import Presentation
        from pptx.enum.chart import XL_CHART_TYPE

        path = os.path.join(self.roots[0], "q3.pptx")
        self.assertEqual(self.deck(path)["slides"], 5)
        prs = Presentation(path)
        slides = list(prs.slides)
        self.assertEqual([s.shapes.title.text for s in slides], ["Q3 review", "Highlights", "Revenue", "Mix", "Team"])
        self.assertEqual(slides[0].placeholders[1].text, "Prepared by Cookie")
        self.assertEqual([p.text for p in slides[1].placeholders[1].text_frame.paragraphs],
                         ["Revenue up 12%", "Two new markets"])
        self.assertEqual(slides[1].notes_slide.notes_text_frame.text, "Say it slowly.")
        charts = [sh.chart for sh in slides[2].shapes if sh.has_chart]
        self.assertEqual(len(charts), 1)
        self.assertEqual(charts[0].chart_type, XL_CHART_TYPE.COLUMN_CLUSTERED)
        self.assertEqual([s.name for s in charts[0].series], ["2025", "2026"])
        self.assertEqual(list(charts[0].series[1].values), [2, 3, 4.5])
        self.assertEqual(list(charts[0].plots[0].categories), ["Q1", "Q2", "Q3"])
        pie = [sh.chart for sh in slides[3].shapes if sh.has_chart][0]
        self.assertEqual(pie.chart_type, XL_CHART_TYPE.PIE)
        self.assertTrue(any(sh.has_text_frame and "Mostly EU" in sh.text_frame.text for sh in slides[3].shapes))
        table = [sh.table for sh in slides[4].shapes if sh.has_table][0]
        self.assertEqual([[c.text for c in r.cells] for r in table.rows], [["Name", "Role"], ["Ann", "Lead"], ["Bo", "Eng"]])

    def test_paths_are_confined_to_the_roots(self):
        inside = self.roots[0]
        os.mkdir(os.path.join(self.root, "outside"))
        os.symlink(os.path.join(self.root, "outside"), os.path.join(inside, "escape"))
        os.symlink(os.path.join(self.root, "outside", "x.pptx"), os.path.join(inside, "link.pptx"))
        for path, msg in [
            (os.path.join(self.root, "outside", "x.pptx"), "outside"),
            (os.path.join(inside, "..", "outside", "x.pptx"), r"'\.\.'"),
            (os.path.join(inside, "escape", "x.pptx"), "outside"),
            (os.path.join(inside, "link.pptx"), "symbolic link"),
            ("relative.pptx", "absolute"),
            (os.path.join(inside, "deck.key"), r"\.pptx"),
            (os.path.join(inside, "missing", "x.pptx"), "does not exist"),
        ]:
            with self.subTest(path=path), self.assertRaisesRegex(ToolError, msg):
                self.deck(path)
        self.assertEqual(os.listdir(os.path.join(self.root, "outside")), [])
        with self.assertRaisesRegex(ToolError, "without --root"):
            server.tool_make_deck({**DECK, "path": os.path.join(inside, "a.pptx")}, [])

    def test_no_overwrite_unless_asked(self):
        path = os.path.join(self.roots[0], "q3.pptx")
        self.deck(path)
        with self.assertRaisesRegex(ToolError, "overwrite: true"):
            self.deck(path)
        self.assertGreater(self.deck(path, overwrite=True)["bytes"], 0)
        self.assertEqual(os.listdir(self.roots[0]), ["q3.pptx"])

    def test_bad_slides_are_refused_before_writing(self):
        bad = [
            [{"title": "x", "chart": {"type": "radar", "categories": ["a"], "series": [{"name": "s", "values": [1]}]}}],
            [{"title": "x", "chart": {"type": "line", "categories": ["a", "b"], "series": [{"name": "s", "values": [1]}]}}],
            [{"title": "x", "chart": {"type": "pie", "categories": ["a"],
                                      "series": [{"name": "s", "values": [1]}, {"name": "t", "values": [2]}]}}],
            [{"title": "x", "table": {"header": ["a", "b"], "rows": [["1"]]}}],
            [{"bullets": ["no title"]}],
            [],
        ]
        for slides in bad:
            with self.subTest(slides=slides), self.assertRaises(ToolError):
                server.tool_make_deck({"title": "t", "slides": slides,
                                       "path": os.path.join(self.roots[0], "bad.pptx")}, self.roots)
        self.assertEqual(os.listdir(self.roots[0]), [])


class Protocol(unittest.TestCase):
    """Spawns server.py and speaks MCP to it the way Cookie's core does (one JSON object per line)."""

    def test_round_trip(self):
        with tempfile.TemporaryDirectory() as root:
            proc = subprocess.Popen([sys.executable, "-I", "-B", os.path.join(HERE, "server.py"), "--root", root],
                                    stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True)
            self.addCleanup(proc.stderr.close)
            self.addCleanup(proc.stdout.close)
            self.addCleanup(proc.kill)

            def send(message):
                proc.stdin.write(json.dumps(message) + "\n")
                proc.stdin.flush()

            def call(msg_id, method, params=None):
                send({"jsonrpc": "2.0", "id": msg_id, "method": method, "params": params or {}})
                reply = json.loads(proc.stdout.readline())
                self.assertEqual(reply["id"], msg_id)
                return reply

            init = call(1, "initialize", {"protocolVersion": "2025-06-18", "capabilities": {},
                                          "clientInfo": {"name": "cookie-harness", "version": "0"}})
            self.assertEqual(init["result"]["protocolVersion"], "2025-06-18")
            self.assertIn("tools", init["result"]["capabilities"])
            send({"jsonrpc": "2.0", "method": "notifications/initialized"})
            self.assertEqual(call(2, "ping")["result"], {})

            tools = call(3, "tools/list")["result"]
            self.assertNotIn("nextCursor", tools)
            self.assertEqual([t["name"] for t in tools["tools"]], ["fetch_page", "fetch_pages", "check_quote", "make_deck"])
            for tool in tools["tools"]:
                self.assertEqual(tool["inputSchema"]["type"], "object")

            path = os.path.join(os.path.realpath(root), "d.pptx")
            made = call(4, "tools/call", {"name": "make_deck", "arguments": {**DECK, "path": path}})["result"]
            self.assertFalse(made["isError"])
            self.assertEqual(json.loads(made["content"][0]["text"])["path"], path)
            self.assertTrue(os.path.exists(path))

            refused = call(5, "tools/call", {"name": "fetch_page", "arguments": {"url": "http://127.0.0.1/"}})["result"]
            self.assertTrue(refused["isError"])
            self.assertIn("non-public", refused["content"][0]["text"])

            self.assertEqual(call(6, "tools/call", {"name": "nope", "arguments": {}})["error"]["code"], -32602)
            self.assertEqual(call(7, "resources/list")["error"]["code"], -32601)
            old = call(8, "initialize", {"protocolVersion": "2024-11-05"})
            self.assertEqual(old["result"]["protocolVersion"], "2024-11-05")

            proc.stdin.write("not json\n")
            proc.stdin.flush()
            self.assertEqual(json.loads(proc.stdout.readline())["error"]["code"], -32700)
            proc.stdin.close()
            self.assertEqual(proc.wait(timeout=10), 0)


if __name__ == "__main__":
    unittest.main()
