"""Cookie's Python tools, served over MCP stdio (newline-delimited JSON-RPC 2.0).

Run with the bundled interpreter, isolated and without writing bytecode into
the signed app:

    vendor/python/bin/python3 -I -B cookie-py/server.py --root ~/Documents/Cookie

Tools: fetch_page, fetch_pages, check_quote, make_deck. See README.md.
"""

import argparse
import hashlib
import http.client
import io
import ipaddress
import json
import os
import re
import socket
import ssl
import sys
import threading
import time
import unicodedata
import zlib
from collections import OrderedDict
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
from urllib.parse import urljoin, urlsplit

SERVER_VERSION = "0.1.0"
# Newest first; the client's version is echoed when we speak it.
PROTOCOL_VERSIONS = ("2025-06-18", "2025-03-26", "2024-11-05")

USER_AGENT = "Cookie/0.1 (+https://enkeon.com)"
MAX_BYTES = 5 * 1024 * 1024
TIMEOUT_S = 20.0
MAX_REDIRECTS = 5
MAX_URLS = 5
DEFAULT_MAX_CHARS = 20000
REDIRECT_CODES = (301, 302, 303, 307, 308)
TEXT_TYPES = ("text/html", "application/xhtml+xml", "text/plain")


class ToolError(Exception):
    """A failure the model should read (returned with isError: true)."""


# ---------------------------------------------------------------- SSRF guard

# Named for the README and the tests; `refusal` also refuses anything else
# that is not globally routable (ipaddress's is_global), e.g. 192.0.0.0/24,
# 198.18.0.0/15, documentation ranges, 240.0.0.0/4.
BLOCKED_NETWORKS = [
    ("unspecified", ipaddress.ip_network("0.0.0.0/8")),
    ("loopback", ipaddress.ip_network("127.0.0.0/8")),
    ("private", ipaddress.ip_network("10.0.0.0/8")),
    ("private", ipaddress.ip_network("172.16.0.0/12")),
    ("private", ipaddress.ip_network("192.168.0.0/16")),
    ("CGNAT", ipaddress.ip_network("100.64.0.0/10")),
    ("link-local / cloud metadata", ipaddress.ip_network("169.254.0.0/16")),
    ("multicast", ipaddress.ip_network("224.0.0.0/4")),
    ("unspecified", ipaddress.ip_network("::/128")),
    ("loopback", ipaddress.ip_network("::1/128")),
    ("unique-local", ipaddress.ip_network("fc00::/7")),
    ("link-local", ipaddress.ip_network("fe80::/10")),
    ("multicast", ipaddress.ip_network("ff00::/8")),
]


def refusal(ip):
    """Why `ip` may not be fetched, or None if it is a public address."""
    ip = ipaddress.ip_address(ip)
    if ip.version == 6:
        # An IPv4 address carried inside IPv6 is judged as that IPv4 address.
        embedded = ip.ipv4_mapped or ip.sixtofour
        if embedded is None and ip in ipaddress.ip_network("64:ff9b::/96"):
            embedded = ipaddress.IPv4Address(int(ip) & 0xFFFFFFFF)
        if embedded is not None:
            why = refusal(embedded)
            return f"{why} (inside {ip})" if why else None
    for name, net in BLOCKED_NETWORKS:
        if ip.version == net.version and ip in net:
            return f"{ip} is {name}"
    if not ip.is_global or ip.is_multicast or ip.is_reserved:
        return f"{ip} is not a public address"
    return None


def system_resolve(host, port):
    infos = socket.getaddrinfo(host, port, type=socket.SOCK_STREAM)
    return [info[4][0] for info in infos]


def vet(url, resolve):
    """Check one hop. Returns (scheme, host, port, ip, target)."""
    try:
        parts = urlsplit(url)
        port = parts.port
    except ValueError as e:
        raise ToolError(f"not a valid URL: {url} ({e})")
    scheme = parts.scheme.lower()
    if scheme not in ("http", "https"):
        raise ToolError(f"only http and https URLs are fetched, not {scheme or 'a relative URL'}: {url}")
    if parts.username is not None or parts.password is not None:
        raise ToolError("URLs with a user name or password are not fetched")
    host = parts.hostname
    if not host:
        raise ToolError(f"the URL has no host: {url}")
    try:
        host = host.encode("idna").decode("ascii")
    except UnicodeError:
        raise ToolError(f"the host name is not valid: {parts.hostname}")
    port = port or (443 if scheme == "https" else 80)
    try:
        addresses = resolve(host, port)
    except (OSError, UnicodeError) as e:
        raise ToolError(f"could not resolve {host}: {e}")
    if not addresses:
        raise ToolError(f"{host} has no address")
    for address in addresses:  # every answer, so a mixed DNS reply cannot slip one through
        why = refusal(address.split("%")[0])
        if why:
            raise ToolError(f"refused {url}: {host} resolves to a non-public address ({why})")
    target = (parts.path or "/") + (f"?{parts.query}" if parts.query else "")
    return scheme, host, port, addresses[0], target


class _PinnedHTTP(http.client.HTTPConnection):
    """Connects to the vetted IP, never re-resolving the name (no DNS rebinding)."""

    def __init__(self, host, port, ip, timeout):
        super().__init__(host, port, timeout=timeout)
        self._ip = ip

    def connect(self):
        self.sock = self.pinned_sock = socket.create_connection((self._ip, self.port), self.timeout)


class _PinnedHTTPS(_PinnedHTTP):
    def connect(self):
        super().connect()
        context = ssl.create_default_context()  # verifies the certificate against the name
        self.sock = self.pinned_sock = context.wrap_socket(self.sock, server_hostname=self.host)


def http_transport(scheme, host, port, ip, target, headers, deadline):
    """One GET with no redirect following. Returns (status, headers, body)."""
    remaining = deadline - time.monotonic()
    if remaining <= 0:
        raise ToolError(f"no answer in {TIMEOUT_S:.0f} s")
    cls = _PinnedHTTPS if scheme == "https" else _PinnedHTTP
    conn = cls(host, port, ip, remaining)
    try:
        conn.request("GET", target, headers=headers)
        resp = conn.getresponse()
        head = {k.lower(): v for k, v in resp.getheaders()}
        if resp.status in REDIRECT_CODES:
            return resp.status, head, b""
        length = head.get("content-length", "")
        if length.isdigit() and int(length) > MAX_BYTES:
            raise ToolError(f"the page is larger than {MAX_BYTES // (1024 * 1024)} MB")
        chunks, total = [], 0
        while True:
            if time.monotonic() > deadline:
                raise ToolError(f"no complete answer in {TIMEOUT_S:.0f} s")
            if resp.isclosed():  # the whole body is in; the socket is already closed
                break
            # conn.sock is None once the response owns the socket, so keep our own handle.
            conn.pinned_sock.settimeout(max(0.1, deadline - time.monotonic()))
            chunk = resp.read(64 * 1024)
            if not chunk:
                break
            total += len(chunk)
            if total > MAX_BYTES:
                raise ToolError(f"the page is larger than {MAX_BYTES // (1024 * 1024)} MB")
            chunks.append(chunk)
        return resp.status, head, b"".join(chunks)
    except (OSError, http.client.HTTPException) as e:
        raise ToolError(f"could not fetch {host}: {e}")
    finally:
        conn.close()


def decompress(body, encoding):
    encoding = (encoding or "identity").strip().lower()
    if encoding in ("", "identity"):
        return body
    if encoding not in ("gzip", "x-gzip", "deflate"):
        raise ToolError(f"the server sent an unsupported encoding: {encoding}")
    wbits = 47 if encoding != "deflate" else 15  # 47 = gzip or zlib header, auto
    d = zlib.decompressobj(wbits)
    try:
        out = d.decompress(body, MAX_BYTES + 1)
    except zlib.error:
        if encoding != "deflate":
            raise ToolError("the compressed page is corrupt")
        out = zlib.decompressobj(-15).decompress(body, MAX_BYTES + 1)  # raw deflate
    if len(out) > MAX_BYTES:
        raise ToolError(f"the page is larger than {MAX_BYTES // (1024 * 1024)} MB uncompressed")
    return out


def fetch_raw(url, resolve=system_resolve, transport=http_transport):
    """GET `url`, following up to 5 redirects by hand, vetting every hop."""
    deadline = time.monotonic() + TIMEOUT_S
    headers = {
        "User-Agent": USER_AGENT,
        "Accept": "text/html,application/xhtml+xml,text/plain;q=0.9,*/*;q=0.1",
        "Accept-Encoding": "gzip, deflate",
        "Connection": "close",
    }
    current = url
    for _ in range(MAX_REDIRECTS + 1):
        scheme, host, port, ip, target = vet(current, resolve)
        status, head, body = transport(scheme, host, port, ip, target, headers, deadline)
        if status in REDIRECT_CODES:
            location = head.get("location")
            if not location:
                raise ToolError(f"HTTP {status} redirect without a Location from {current}")
            current = urljoin(current, location.strip())
            continue
        if not 200 <= status < 300:
            raise ToolError(f"HTTP {status} from {current}")
        return current, head, decompress(body, head.get("content-encoding"))
    raise ToolError(f"more than {MAX_REDIRECTS} redirects starting at {url}")


# ------------------------------------------------------------- extraction

_BLOCK_TAG = re.compile(
    r"<(/?)(p|div|br|hr|li|ul|ol|dl|dt|dd|h[1-6]|tr|td|th|table|section|article|main|header|footer|"
    r"aside|nav|blockquote|pre|figure|figcaption|address|details|summary)\b",
    re.IGNORECASE,
)
_META_CHARSET = re.compile(rb"""<meta[^>]+charset\s*=\s*["']?\s*([A-Za-z0-9_\-:.]+)""", re.IGNORECASE)
IGNORED_TAGS = ("script", "style", "noscript", "template", "svg", "head", "iframe", "object", "canvas", "button", "select")
NOT_CONTENT = IGNORED_TAGS + ("nav", "footer", "header", "aside", "form")


def decode(body, content_type):
    charset = None
    m = re.search(r"charset\s*=\s*[\"']?([A-Za-z0-9_\-:.]+)", content_type or "", re.IGNORECASE)
    if m:
        charset = m.group(1)
    else:
        m = _META_CHARSET.search(body[:4096])
        charset = m.group(1).decode("ascii") if m else None
    try:
        return body.decode(charset or "utf-8", errors="replace")
    except LookupError:
        return body.decode("utf-8", errors="replace")


def tidy(text):
    lines = (re.sub(r"[^\S\n]+", " ", line).strip() for line in text.split("\n"))
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()


def extract(html, url):
    """(title, text) using Scrapling's parser: the main region if there is one, else all visible text."""
    from scrapling.parser import Selector

    # A newline before each block tag, so block boundaries survive while inline
    # elements ("<b>Hello</b> world") keep the document's own spacing.
    page = Selector(_BLOCK_TAG.sub(r"\n<\1\2", html), url=url)

    def visible(node, ignore):
        return tidy(str(node.get_all_text(separator="", strip=False, ignore_tags=ignore, valid_values=False)))

    title = ""
    for query in ("title::text", 'meta[property="og:title"]::attr(content)', "h1"):
        found = page.css(query)
        if found:
            value = found[0] if query == "h1" else found.get()
            title = visible(value, IGNORED_TAGS) if query == "h1" else tidy(str(value or ""))
            if title:
                break

    best = ""
    for query in ("article", "main", '[role="main"]'):
        for node in page.css(query):
            text = visible(node, NOT_CONTENT)
            if len(text) > len(best):
                best = text
    if len(best) < 200:  # no real main region: fall back to everything visible
        body = page.css("body")
        best = visible(body[0] if body else page, IGNORED_TAGS)
    return title, best


# Pages fetched by this process, so check_quote checks the very text the model was given.
_CACHE = OrderedDict()
_CACHE_MAX = 64
_CACHE_LOCK = threading.Lock()


def load_page(url, resolve=system_resolve, transport=http_transport):
    if not isinstance(url, str) or not url.strip():
        raise ToolError("url must be a non-empty string")
    url = url.strip()
    final_url, head, body = fetch_raw(url, resolve, transport)
    content_type = head.get("content-type", "")
    mime = content_type.split(";")[0].strip().lower()
    if mime and mime not in TEXT_TYPES:
        raise ToolError(f"{final_url} is {mime}, not a web page")
    raw = decode(body, content_type)
    if mime == "text/plain":
        title, text = "", tidy(raw)
    else:
        title, text = extract(raw, final_url)
    page = {
        "url": url,
        "final_url": final_url,
        "title": title,
        "fetched_at": datetime.now(timezone.utc).isoformat(timespec="seconds").replace("+00:00", "Z"),
        "sha256": hashlib.sha256(text.encode("utf-8")).hexdigest(),
        "text": text,
    }
    with _CACHE_LOCK:
        _CACHE[url] = page
        _CACHE.move_to_end(url)
        while len(_CACHE) > _CACHE_MAX:
            _CACHE.popitem(last=False)
    return page


def present(page, max_chars):
    text = page["text"]
    out = {k: page[k] for k in ("url", "final_url", "title", "fetched_at", "sha256")}
    out["chars"] = len(text)
    out["truncated"] = len(text) > max_chars
    out["text"] = text[:max_chars]
    if out["truncated"]:
        out["note"] = f"text truncated to the first {max_chars} of {len(text)} characters; sha256 is of the full text"
    return out


def _max_chars(args):
    value = args.get("max_chars", DEFAULT_MAX_CHARS)
    if not isinstance(value, int) or isinstance(value, bool) or not 1 <= value <= 200000:
        raise ToolError("max_chars must be an integer from 1 to 200000")
    return value


def tool_fetch_page(args, resolve=system_resolve, transport=http_transport):
    return present(load_page(args.get("url"), resolve, transport), _max_chars(args))


def tool_fetch_pages(args, resolve=system_resolve, transport=http_transport):
    urls = args.get("urls")
    if not isinstance(urls, list) or not urls or len(urls) > MAX_URLS:
        raise ToolError(f"urls must be a list of 1 to {MAX_URLS} URLs")
    max_chars = _max_chars(args)

    def one(url):
        try:
            return {"ok": True, **present(load_page(url, resolve, transport), max_chars)}
        except ToolError as e:
            return {"ok": False, "url": url, "error": str(e)}

    with ThreadPoolExecutor(max_workers=len(urls)) as pool:
        return {"results": list(pool.map(one, urls))}


def normalise(text):
    return re.sub(r"\s+", " ", unicodedata.normalize("NFC", text)).strip()


_SENTENCE_END = re.compile(r"[.!?](?=\s)")


def find_quote(text, quote):
    """Where `quote` occurs in `text` (whitespace-normalised, case-sensitive) and its sentence."""
    hay, needle = normalise(text), normalise(quote)
    if not needle:
        raise ToolError("quote must contain some text")
    at = hay.find(needle)
    if at < 0:
        return {"found": False, "quote": needle, "occurrences": 0, "sentence": None}
    ends = [m.end() for m in _SENTENCE_END.finditer(hay, 0, at)]
    start = max(ends[-1] if ends else 0, at - 300)
    after = _SENTENCE_END.search(hay, at + len(needle))
    stop = min(after.end() if after else len(hay), at + len(needle) + 300)
    return {"found": True, "quote": needle, "occurrences": hay.count(needle), "sentence": hay[start:stop].strip()}


def tool_check_quote(args, resolve=system_resolve, transport=http_transport):
    url, quote = args.get("url"), args.get("quote")
    if not isinstance(quote, str):
        raise ToolError("quote must be a string")
    with _CACHE_LOCK:
        page = _CACHE.get(url.strip()) if isinstance(url, str) else None
    from_cache = page is not None
    if page is None:
        page = load_page(url, resolve, transport)
    result = find_quote(page["text"], quote)
    return {
        "url": page["url"],
        "final_url": page["final_url"],
        "fetched_at": page["fetched_at"],
        "sha256": page["sha256"],
        "from_cache": from_cache,
        **result,
    }


# ------------------------------------------------------------------ decks

ROOTS = []


def confine(path, roots, overwrite):
    """The real path to write, inside one of `roots`, or ToolError."""
    if not roots:
        raise ToolError("no folder is open for writing (the server was started without --root)")
    if not isinstance(path, str) or not path:
        raise ToolError("path must be a non-empty string")
    if not os.path.isabs(path):
        raise ToolError("path must be absolute")
    if ".." in path.replace("\\", "/").split("/"):
        raise ToolError("path may not contain '..'")
    if not path.lower().endswith(".pptx"):
        raise ToolError("path must end in .pptx")
    parent = os.path.dirname(path)
    if not os.path.isdir(parent):
        raise ToolError(f"the folder {parent} does not exist")
    target = os.path.join(os.path.realpath(parent), os.path.basename(path))
    if not any(os.path.commonpath([target, root]) == root for root in roots):
        raise ToolError(f"{path} is outside the folders Cookie may write to")
    if os.path.islink(target):
        raise ToolError(f"{path} is a symbolic link")
    if os.path.lexists(target) and not overwrite:
        raise ToolError(f"{path} exists; pass overwrite: true to replace it")
    if os.path.isdir(target):
        raise ToolError(f"{path} is a folder")
    return target


def _strings(value, what, limit):
    if not isinstance(value, list) or len(value) > limit or not all(isinstance(v, str) for v in value):
        raise ToolError(f"{what} must be a list of at most {limit} strings")
    return value


def _check_chart(chart, n):
    if not isinstance(chart, dict) or chart.get("type") not in ("column", "bar", "line", "pie"):
        raise ToolError(f"slide {n}: chart.type must be column, bar, line or pie")
    categories = _strings(chart.get("categories"), f"slide {n}: chart.categories", 100)
    series = chart.get("series")
    if not isinstance(series, list) or not series or len(series) > 20:
        raise ToolError(f"slide {n}: chart.series must be a list of 1 to 20 series")
    if chart["type"] == "pie" and len(series) != 1:
        raise ToolError(f"slide {n}: a pie chart takes exactly one series")
    for s in series:
        values = s.get("values") if isinstance(s, dict) else None
        if not isinstance(s, dict) or not isinstance(s.get("name"), str):
            raise ToolError(f"slide {n}: each series needs a name")
        if not isinstance(values, list) or len(values) != len(categories) or not all(
            isinstance(v, (int, float)) and not isinstance(v, bool) for v in values
        ):
            raise ToolError(f"slide {n}: series '{s['name']}' needs one number per category")


def _check_table(table, n):
    header = _strings(table.get("header") if isinstance(table, dict) else None, f"slide {n}: table.header", 12)
    rows = table.get("rows")
    if not header or not isinstance(rows, list) or len(rows) > 50:
        raise ToolError(f"slide {n}: a table needs a header and at most 50 rows")
    for row in rows:
        if not isinstance(row, list) or len(row) != len(header):
            raise ToolError(f"slide {n}: every table row needs {len(header)} cells")


def build_deck(title, subtitle, slides):
    """The .pptx bytes. Charts are native PowerPoint charts (editable data)."""
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.util import Inches, Pt

    kinds = {
        "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
        "bar": XL_CHART_TYPE.BAR_CLUSTERED,
        "line": XL_CHART_TYPE.LINE_MARKERS,
        "pie": XL_CHART_TYPE.PIE,
    }
    prs = Presentation()  # the default template: 0 Title, 1 Title and Content, 5 Title Only
    # 16:9, as decks are shown today. The default template is 4:3 (10 in wide), so its
    # placeholders are stretched sideways with the slide; heights are already 7.5 in.
    stretch = Inches(13.333) / prs.slide_width
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for owner in [prs.slide_master, *prs.slide_layouts]:
        for placeholder in owner.placeholders:
            if placeholder.left is not None and placeholder.width is not None:
                placeholder.left, placeholder.width = int(placeholder.left * stretch), int(placeholder.width * stretch)
    first = prs.slides.add_slide(prs.slide_layouts[0])
    first.shapes.title.text = title
    if subtitle:
        first.placeholders[1].text = subtitle
    else:
        first.placeholders[1]._element.getparent().remove(first.placeholders[1]._element)

    width, height = prs.slide_width, prs.slide_height
    margin, top = Inches(0.5), Inches(1.6)
    for spec in slides:
        bullets, chart, table = spec.get("bullets") or [], spec.get("chart"), spec.get("table")
        visual = chart or table
        slide = prs.slides.add_slide(prs.slide_layouts[5 if visual else 1])
        slide.shapes.title.text = spec["title"]
        if not visual:
            body = slide.placeholders[1].text_frame
            for i, line in enumerate(bullets):
                para = body.paragraphs[0] if i == 0 else body.add_paragraph()
                para.text = line
        else:
            left, avail = margin, width - 2 * margin
            if bullets:  # bullets in a left column, the chart or table beside them
                col = int(avail * 0.4)
                box = slide.shapes.add_textbox(left, top, col, height - top - margin).text_frame
                box.word_wrap = True
                for i, line in enumerate(bullets):
                    para = box.paragraphs[0] if i == 0 else box.add_paragraph()
                    para.text = f"• {line}"
                    para.font.size = Pt(16)
                left, avail = left + col + Inches(0.2), avail - col - Inches(0.2)
            area = (left, top, avail, height - top - margin)
            if chart:
                data = CategoryChartData()
                data.categories = chart["categories"]
                for s in chart["series"]:
                    data.add_series(s["name"], s["values"])
                graphic = slide.shapes.add_chart(kinds[chart["type"]], *area, data).chart
                if chart["type"] == "pie" or len(chart["series"]) > 1:
                    graphic.has_legend = True
                    graphic.legend.position = XL_LEGEND_POSITION.BOTTOM
                    graphic.legend.include_in_layout = False
                if chart["type"] == "pie":
                    graphic.plots[0].has_data_labels = True
                    graphic.plots[0].data_labels.show_percentage = True
                    graphic.plots[0].data_labels.number_format = "0%"
                    graphic.plots[0].data_labels.number_format_is_linked = False
            if table:
                rows = [table["header"]] + table["rows"]
                shape = slide.shapes.add_table(len(rows), len(table["header"]), *area[:3], Inches(0.4) * len(rows))
                for r, row in enumerate(rows):
                    for c, cell in enumerate(row):
                        shape.table.cell(r, c).text = "" if cell is None else str(cell)
        if spec.get("notes"):
            slide.notes_slide.notes_text_frame.text = spec["notes"]
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def tool_make_deck(args, roots=None):
    roots = ROOTS if roots is None else roots
    title, subtitle, slides = args.get("title"), args.get("subtitle"), args.get("slides")
    overwrite = args.get("overwrite", False) is True
    if not isinstance(title, str) or not title.strip():
        raise ToolError("title must be a non-empty string")
    if subtitle is not None and not isinstance(subtitle, str):
        raise ToolError("subtitle must be a string")
    if not isinstance(slides, list) or not slides or len(slides) > 100:
        raise ToolError("slides must be a list of 1 to 100 slides")
    for n, spec in enumerate(slides, 1):
        if not isinstance(spec, dict) or not isinstance(spec.get("title"), str):
            raise ToolError(f"slide {n}: needs a title")
        _strings(spec.get("bullets") or [], f"slide {n}: bullets", 30)
        if spec.get("notes") is not None and not isinstance(spec["notes"], str):
            raise ToolError(f"slide {n}: notes must be a string")
        if spec.get("chart") is not None:
            _check_chart(spec["chart"], n)
        if spec.get("table") is not None:
            _check_table(spec["table"], n)
        if spec.get("chart") is not None and spec.get("table") is not None:
            raise ToolError(f"slide {n}: a slide holds a chart or a table, not both")
    target = confine(args.get("path"), roots, overwrite)
    data = build_deck(title, subtitle, slides)
    if overwrite:
        tmp = f"{target}.{os.getpid()}.tmp"
        fd = os.open(tmp, os.O_WRONLY | os.O_CREAT | os.O_EXCL | os.O_NOFOLLOW, 0o644)
        with os.fdopen(fd, "wb") as f:
            f.write(data)
        os.replace(tmp, target)  # replaces a name, never writes through a link
    else:
        try:  # O_EXCL: never clobbers a file that appeared since the check
            fd = os.open(target, os.O_WRONLY | os.O_CREAT | os.O_EXCL | os.O_NOFOLLOW, 0o644)
        except FileExistsError:
            raise ToolError(f"{target} exists; pass overwrite: true to replace it")
        with os.fdopen(fd, "wb") as f:
            f.write(data)
    return {"path": target, "slides": len(slides) + 1, "bytes": len(data)}


# -------------------------------------------------------------------- MCP

_URL = {"type": "string", "description": "An http or https URL."}
_MAX_CHARS = {"type": "integer", "minimum": 1, "maximum": 200000, "default": DEFAULT_MAX_CHARS,
              "description": "Most characters of text to return."}
_SERIES = {"type": "object", "required": ["name", "values"], "properties": {
    "name": {"type": "string"}, "values": {"type": "array", "items": {"type": "number"}}}}

TOOLS = [
    {
        "name": "fetch_page",
        "description": "Fetch one public web page (static HTML, no JavaScript) and return its title and readable text, "
                       "with the final URL after redirects, the UTC fetch time and a SHA-256 of the text.",
        "inputSchema": {"type": "object", "required": ["url"],
                        "properties": {"url": _URL, "max_chars": _MAX_CHARS}},
        "annotations": {"title": "Fetch a web page", "readOnlyHint": True, "openWorldHint": True},
    },
    {
        "name": "fetch_pages",
        "description": f"Fetch up to {MAX_URLS} public web pages at once; each succeeds or fails on its own.",
        "inputSchema": {"type": "object", "required": ["urls"], "properties": {
            "urls": {"type": "array", "items": _URL, "minItems": 1, "maxItems": MAX_URLS},
            "max_chars": _MAX_CHARS}},
        "annotations": {"title": "Fetch web pages", "readOnlyHint": True, "openWorldHint": True},
    },
    {
        "name": "check_quote",
        "description": "Check that a quote appears word for word on a page (whitespace-normalised, case-sensitive) "
                       "and return the sentence around it. Uses the text fetch_page returned earlier in this session "
                       "when there is one, else fetches the page.",
        "inputSchema": {"type": "object", "required": ["url", "quote"],
                        "properties": {"url": _URL, "quote": {"type": "string"}}},
        "annotations": {"title": "Check a quote", "readOnlyHint": True, "openWorldHint": True},
    },
    {
        "name": "make_deck",
        "description": "Write a PowerPoint (.pptx) deck with a title slide and the given slides; charts are native "
                       "PowerPoint charts. The path must be inside a folder Cookie may write to; an existing file is "
                       "replaced only with overwrite: true.",
        "inputSchema": {"type": "object", "required": ["path", "title", "slides"], "properties": {
            "path": {"type": "string", "description": "Absolute path ending in .pptx."},
            "title": {"type": "string"},
            "subtitle": {"type": "string"},
            "overwrite": {"type": "boolean", "default": False},
            "slides": {"type": "array", "minItems": 1, "maxItems": 100, "items": {
                "type": "object", "required": ["title"], "properties": {
                    "title": {"type": "string"},
                    "bullets": {"type": "array", "items": {"type": "string"}},
                    "notes": {"type": "string"},
                    "chart": {"type": "object", "required": ["type", "categories", "series"], "properties": {
                        "type": {"enum": ["column", "bar", "line", "pie"]},
                        "categories": {"type": "array", "items": {"type": "string"}},
                        "series": {"type": "array", "items": _SERIES}}},
                    "table": {"type": "object", "required": ["header", "rows"], "properties": {
                        "header": {"type": "array", "items": {"type": "string"}},
                        "rows": {"type": "array", "items": {"type": "array"}}}}}}}}},
        "annotations": {"title": "Make a slide deck", "readOnlyHint": False, "destructiveHint": True,
                        "idempotentHint": False, "openWorldHint": False},
    },
]

HANDLERS = {
    "fetch_page": tool_fetch_page,
    "fetch_pages": tool_fetch_pages,
    "check_quote": tool_check_quote,
    "make_deck": tool_make_deck,
}


def handle(message):
    """One JSON-RPC message in, the reply (or None for a notification) out."""
    if not isinstance(message, dict) or message.get("jsonrpc") != "2.0" or not isinstance(message.get("method"), str):
        return {"jsonrpc": "2.0", "id": message.get("id") if isinstance(message, dict) else None,
                "error": {"code": -32600, "message": "invalid request"}}
    method, params, msg_id = message["method"], message.get("params") or {}, message.get("id")
    if "id" not in message:
        return None  # notifications (initialized, cancelled) need no answer

    def ok(result):
        return {"jsonrpc": "2.0", "id": msg_id, "result": result}

    def err(code, text):
        return {"jsonrpc": "2.0", "id": msg_id, "error": {"code": code, "message": text}}

    if method == "initialize":
        asked = params.get("protocolVersion") if isinstance(params, dict) else None
        return ok({
            "protocolVersion": asked if asked in PROTOCOL_VERSIONS else PROTOCOL_VERSIONS[0],
            "capabilities": {"tools": {"listChanged": False}},
            "serverInfo": {"name": "cookie-py", "version": SERVER_VERSION},
        })
    if method == "ping":
        return ok({})
    if method == "tools/list":
        return ok({"tools": TOOLS})
    if method == "tools/call":
        name = params.get("name") if isinstance(params, dict) else None
        arguments = params.get("arguments") or {} if isinstance(params, dict) else {}
        if name not in HANDLERS:
            return err(-32602, f"unknown tool: {name}")
        if not isinstance(arguments, dict):
            return err(-32602, "arguments must be an object")
        try:
            result, is_error = HANDLERS[name](arguments), False
            text = json.dumps(result, ensure_ascii=False)
        except ToolError as e:
            text, is_error = str(e), True
        except Exception as e:  # a bug here must not take the server down
            print(f"cookie-py: {name} failed: {e!r}", file=sys.stderr, flush=True)
            text, is_error = f"{name} failed: {type(e).__name__}: {e}", True
        return ok({"content": [{"type": "text", "text": text}], "isError": is_error})
    return err(-32601, f"method not found: {method}")


def serve(stdin, stdout):
    for line in stdin:
        if not line.strip():
            continue
        try:
            message = json.loads(line)
        except ValueError:
            reply = {"jsonrpc": "2.0", "id": None, "error": {"code": -32700, "message": "parse error"}}
        else:
            reply = handle(message)
        if reply is not None:
            stdout.write(json.dumps(reply, ensure_ascii=False) + "\n")
            stdout.flush()


def main(argv=None):
    parser = argparse.ArgumentParser(description="Cookie's Python tools over MCP stdio.")
    parser.add_argument("--root", action="append", default=[], help="a folder make_deck may write in (repeatable)")
    args = parser.parse_args(argv)
    for root in args.root:
        real = os.path.realpath(os.path.expanduser(root))
        if not os.path.isdir(real):
            parser.error(f"--root {root} is not a folder")
        ROOTS.append(real)
    # stdout carries protocol only; anything a library prints goes to stderr.
    out = sys.stdout
    sys.stdout = sys.stderr
    serve(sys.stdin, out)


if __name__ == "__main__":
    main()
